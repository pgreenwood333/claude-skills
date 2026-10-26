#!/usr/bin/env python3
"""Build the QA/QC Excel workbook from the reviewed findings file.

  build_report.py <findings_reviewed.json> <index_dir> <out.xlsx>

Tabs (only for the tasks that ran): Cover, Summary Dashboard, Index Check, Spelling, Overlapping Text,
Table 1 - EQ vs SQS, Table 2 - EQ Recap vs EQ, Discrepancy Detail, and Rejected (flags the orchestrator
dismissed on review, kept for transparency). Run markup_pdf.py first so the Marked Up column is filled.
"""
import datetime, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from qaqc_common import load_json, TASKS

# muted fills (status colours from the MDOT reconciliation convention)
FILL = {
    "MATCH": "D9EAD3", "MISMATCH": "F4CCCC", "VERIFIED WITH FOOTNOTE": "FFF2CC",
    "ESTIMATED-NOT RECONCILED": "F8D7E3", "N/A-NO RECAP EXPECTED": "E7E6E6", "NOT FOUND ON EQ": None,
    "NOT FOUND ON SQS": None, "RECAP NOT FOUND": "FCE4D0", "PARTIAL-SOME RECAP SHEETS MISSING": "E4DFEC",
    "high": "F4CCCC", "medium": "FFF2CC", "low": None,
}
HEAD = PatternFill("solid", fgColor="3F4E5C")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(wb, title, headers, rows, widths, status_col=None, num_cols=()):
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
    for r in rows:
        ws.append(r)
        row = ws[ws.max_row]
        key = r[status_col] if status_col is not None else None
        fill = FILL.get(key) if key else None
        for i, c in enumerate(row):
            c.alignment = WRAP
            c.border = BORDER
            if fill:
                c.fill = PatternFill("solid", fgColor=fill)
            if i in num_cols and isinstance(c.value, (int, float)):
                c.number_format = "#,##0.###"
                c.alignment = Alignment(horizontal="right", vertical="top")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def jl(v):
    return ", ".join(str(x) for x in v) if isinstance(v, (list, tuple)) else ("" if v is None else str(v))


def main():
    data = load_json(sys.argv[1])
    idx = load_json(os.path.join(sys.argv[2], "index.json"))
    out = sys.argv[3]
    tasks = data.get("tasks", list(TASKS))
    F = data["findings"]
    kept = [f for f in F if (f.get("review") or {}).get("verdict") != "rejected"]
    rejected = [f for f in F if (f.get("review") or {}).get("verdict") == "rejected"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Cover"
    proj = data.get("project", {})
    info = [("Plan set", idx["summary"]["pdf"]), ("Pages", idx["summary"]["pages"]),
            ("Project", proj.get("project_no", "")), ("County", proj.get("county", "")),
            ("Checks run", ", ".join(TASKS[t] for t in tasks)),
            ("Run date", datetime.date.today().isoformat()),
            ("OCR engine", idx["summary"].get("ocr_engine")),
            ("Pages read by OCR", jl(idx["summary"].get("raster_pages"))),
            ("Note", "Every item is a flag for engineer review, not a finding of fact. Sheet IDs are the keys; "
                     "PDF page numbers are given to help find them. Red highlights in the marked-up PDF match the IDs here.")]
    for k, v in info:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.cell(ws.max_row, 2).alignment = WRAP
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110

    # dashboard
    dash = []
    for t in tasks:
        ft = [f for f in kept if f.get("task") == t]
        dash.append([TASKS[t], sum(f.get("severity") == "high" for f in ft), sum(f.get("severity") == "medium" for f in ft),
                     sum(f.get("severity") == "low" for f in ft), len(ft),
                     sum(f.get("task") == t for f in rejected)])
    q = data.get("quantities") or {}
    if "quantity" in tasks and q:
        for tbl in ("table1", "table2"):
            for st, n in (q.get("counts", {}).get(tbl) or {}).items():
                dash.append([f"{'Table 1' if tbl == 'table1' else 'Table 2'}: {st}", "", "", "", n, ""])
    sheet(wb, "Summary Dashboard", ["Check", "High", "Medium", "Low", "Total", "Rejected on review"], dash,
          [44, 10, 10, 10, 10, 18], num_cols=(1, 2, 3, 4, 5))

    cols = ["ID", "Severity", "Sheet ID(s)", "PDF page(s)", "Finding", "Evidence (as printed)", "Suggested fix",
            "Review", "Review note", "Marked up"]
    widths = [12, 10, 16, 11, 60, 40, 36, 12, 36, 12]

    def rows_for(task, items):
        return [[f["id"], f.get("severity"), jl(f.get("sheet_ids")), jl(f.get("pdf_pages")), f.get("finding"),
                 f.get("evidence"), f.get("suggested_fix"), (f.get("review") or {}).get("verdict"),
                 (f.get("review") or {}).get("note"), f.get("markup")]
                for f in sorted(items, key=lambda f: ("high", "medium", "low").index(f.get("severity", "low")))
                if f.get("task") == task]

    if "index" in tasks:
        sheet(wb, "Index Check", cols, rows_for("index", kept), widths, status_col=1)
    if "spelling" in tasks:
        sheet(wb, "Spelling", cols, rows_for("spelling", kept), widths, status_col=1)
    if "overlap" in tasks:
        sheet(wb, "Overlapping Text", cols, rows_for("overlap", kept), widths, status_col=1)
    if "quantity" in tasks and q:
        t1 = [[i, r.get("section", ""), r["pay_item"], r["description"], r["unit"], r["eq_qty"], jl(r["eq_sheets"]), r["sqs_qty"],
               jl(r["sqs_sheets"]), r["status"], r["notes"]] for i, r in enumerate(q.get("table1", []), 1)]
        sheet(wb, "Table 1 - EQ vs SQS", ["No.", "Bid Schedule", "Pay Item No.", "Description", "Unit", "EQ Qty",
                                          "EQ Sheet ID", "SQS Qty", "SQS Sheet ID", "Status", "Notes"], t1,
              [6, 14, 16, 48, 8, 12, 14, 12, 14, 26, 48], status_col=9, num_cols=(5, 7))
        t2 = [[i, r.get("section", ""), r["description"], r["unit"], r["recap_sum"], jl(r["recap_sheets"]), r["eq_total"], jl(r["eq_sheets"]),
               r["status"], r["notes"]] for i, r in enumerate(q.get("table2", []), 1)]
        sheet(wb, "Table 2 - EQ Recap vs EQ", ["No.", "Bid Schedule", "Description", "Unit", "EQ Recap Sum",
                                               "EQ Recap Sheet IDs", "EQ Total", "EQ Sheet ID", "Status", "Notes"], t2,
              [6, 14, 48, 8, 14, 22, 12, 14, 30, 60], status_col=8, num_cols=(4, 6))
        dd = [[i, d["table"], d.get("pay_item"), d.get("description"), d["values"], d.get("cause")]
              for i, d in enumerate(q.get("discrepancies", []), 1)]
        sheet(wb, "Discrepancy Detail", ["No.", "Table", "Pay Item No.", "Description", "Conflicting values (Sheet IDs)",
                                         "Most likely cause"], dd, [6, 9, 16, 44, 56, 56])
    if rejected:
        sheet(wb, "Rejected", cols, [[f["id"], f.get("severity"), jl(f.get("sheet_ids")), jl(f.get("pdf_pages")),
                                      f.get("finding"), f.get("evidence"), f.get("suggested_fix"), "rejected",
                                      (f.get("review") or {}).get("note"), f.get("markup")] for f in rejected], widths)
    wb.save(out)
    print(f"wrote {out}: {len(kept)} flags, {len(rejected)} rejected")


if __name__ == "__main__":
    main()
