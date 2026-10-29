#!/usr/bin/env python3
"""Build an MDOT Opinion of Probable Construction Cost workbook.

Input: the JSON from read_sq.py. Prices come from mdot_bid_prices.lookup().
Layout follows the firm's OPCC sheet (one tab per bid section):
  C PAY ITEM NO. | D PAY ITEM | E UNIT | F QTY | G UNIT PRICE | H TOTAL | I % | J REMARKS
  L YEAR | M INITIAL COST | N COST IN 2020 | O ESTIMATE YR   (inflation, Q4:R6 rates)
  T..AC  MDOT bid data used (range, price year, quantities, low/avg/high, suggested, method)

Usage:
  build_opcc.py sq.json --project "Lakeland Drive Pedestrian Bridge" --fa "STBG-9999-05(413)" \
      --by GREENWOOD --out "123456-301000-2026-10-29-OPCC.xlsx" [--today 2026-10-29] [--refresh]
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import math
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
import mdot_bid_prices as mbp  # noqa: E402

PERCENT_ITEMS = {  # pay item -> label used in REMARKS
    "201-B001": "Clearing and Grubbing",
    "907-618-A001": "Maintenance of Traffic",
    "620-A001": "Mobilization",
    "699-A001": "Roadway Construction Stakes",
}
DEFAULT_PCT = 0.05
WIDE_RANGE = 10  # flag when high price > 10x low price
UNITS = {
    "ACRE": "Acre", "AC": "Acre", "SY": "Square Yard", "LF": "Linear Feet", "CY": "Cubic Yard", "TON": "Ton",
    "EA": "Each", "LS": "Lump Sum", "GAL": "Gallon", "LBS": "Pounds", "LB": "Pounds", "SF": "Square Feet",
    "HRS": "Hours", "HR": "Hours", "STA": "Station", "MI": "Mile", "MILE": "Mile", "KGAL": "Thousand Gallon",
    "MGAL": "Thousand Gallon",
}

# ------------------------------------------------------------------ pricing
def suggest(item: dict, rec: dict | None) -> dict:
    """Return {'price', 'year', 'method', 'flag'} for one SQ line."""
    code, qty = item["item"], item["qty"]
    if code in PERCENT_ITEMS:
        return {"price": None, "year": None, "method": "% of total items cost", "flag": None, "percent": True}
    if not rec or not rec.get("found"):
        return {"price": None, "year": None, "method": "Flagged", "flag": "No MDOT bid price found"}
    sq_unit = UNITS.get(item["unit"].upper().replace(".", ""), item["unit"])
    if sq_unit.lower() != str(rec["unit"]).lower():
        return {"price": None, "year": None, "method": "Flagged",
                "flag": f"Unit differs: SQ {item['unit']} vs MDOT {rec['unit']}"}
    lo, hi, avg = rec["low_price"], rec["high_price"], rec["avg_price"]
    qmin, qmax = rec["min_qty"], rec["max_qty"]
    year = rec["price_year"]
    if "-PP" in code.upper():
        return {"price": None, "year": None, "method": "Flagged", "flag": "Per Plans item - use judgment"}
    if qmin == qmax and lo == hi:
        return {"price": round(avg, 2), "year": year, "method": "Only one bid price",
                "flag": None, "note": "CHECK: only one MDOT bid price"}
    if lo and hi > WIDE_RANGE * lo:
        # Penny / unbalanced bids make the low end meaningless, so use the average instead of the qty line.
        return {"price": round(avg, 2), "year": year, "method": f"Wide range ({hi / lo:,.0f}x) -> avg",
                "flag": None, "note": f"CHECK: wide MDOT price range (${lo:,.2f} - ${hi:,.2f}), used avg"}
    if qty is None:
        return {"price": None, "year": None, "method": "Flagged", "flag": "No quantity on SQ"}
    if qmin == qmax:
        return {"price": round(avg, 2), "year": year, "method": "One bid quantity -> avg", "flag": None}
    if qty <= qmin:
        return {"price": round(hi, 2), "year": year, "method": "At/below min qty -> high", "flag": None}
    if qty >= qmax:
        return {"price": round(lo, 2), "year": year, "method": "At/above max qty -> low", "flag": None}
    t = (math.log(qty) - math.log(qmin)) / (math.log(qmax) - math.log(qmin))
    return {"price": round(hi + t * (lo - hi), 2), "year": year, "method": "Qty line (log)", "flag": None}


# ------------------------------------------------------------------ workbook
THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # cells the engineer is expected to fill
DATA_FILL = PatternFill("solid", fgColor="EDF2F8")    # MDOT data block
MONEY = '"$"#,##0.00'
MONEY0 = '"$"#,##0'
QTYF = '#,##0.##'
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_sheet(wb, sec, prices, meta, pct_ref):
    ws = wb.create_sheet(sec["name"].title()[:31])
    est_year = meta["today"].year
    ws["D1"] = "OPINION OF PROBABLE CONSTRUCTION COST"; ws["D1"].font = Font(bold=True, size=14)
    ws["D2"] = meta["project"].upper(); ws["E2"] = "DATE:"; ws["G2"] = meta["today"]
    ws["G2"].number_format = "m/d/yyyy"
    ws["D3"] = f"{meta['county']} COUNTY" if meta["county"] else ""
    ws["E3"] = "UPDATED BY:"; ws["G3"] = meta["by"]
    ws["L3"] = "ENTER YEAR & COST"
    ws["D4"] = f"{meta['job_slash']}  {meta['fa']}".strip()
    ws["E4"] = "TOTAL ITEMS COST:"
    ws["Q4"] = "UNIT PRICE INFLATION"
    heads = {"C": "PAY ITEM NO.", "D": "PAY ITEM", "E": "UNIT", "F": "QTY", "G": "UNIT PRICE", "H": "TOTAL",
             "I": "% OF ITEMS", "J": "REMARKS", "L": "INFLATION", "N": "COST IN YR", "O": "ESTIMATE YR",
             "Q": "2020 - 2025 @", "R": 0.08}
    for c, v in heads.items():
        ws[f"{c}5"] = v
    ws["L6"], ws["M6"], ws["N6"], ws["O6"] = "YEAR", "INITIAL COST", 2020, est_year
    ws["Q6"], ws["R6"] = "BEFORE 2020 @", 0.03
    ws["R5"].number_format = ws["R6"].number_format = "0%"
    # MDOT block
    ws["T3"] = (f"MDOT ITEM BID PRICES - aashtoreports.mdot.ms.gov BidPrices report, downloaded "
                f"{meta['downloaded']}. Exact values as published by MDOT.")
    ws["T3"].font = Font(italic=True, size=9)
    mdot_heads = ["SEARCH RANGE", "PRICE YEAR", "MIN QTY", "MAX QTY", "AVG QTY", "LOW PRICE", "AVG PRICE",
                  "HIGH PRICE", "SUGGESTED (BEFORE INFLATION)", "METHOD"]
    ws["T4"] = "MDOT BID DATA USED"; ws["T4"].font = BOLD
    for i, h in enumerate(mdot_heads):
        ws.cell(5, 20 + i, h)
    for row in (5, 6):
        for col in list(range(3, 11)) + list(range(12, 16)) + list(range(20, 30)):
            c = ws.cell(row, col); c.font = BOLD; c.alignment = CENTER; c.fill = HEAD_FILL; c.border = BOX
    ws.merge_cells("L5:M5")

    first = 7
    r = first
    pct_rows = []
    flagged = []
    checks = []
    for it in sec["items"]:
        code = it["item"]
        rec = prices.get(code)
        s = suggest(it, rec)
        ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"], ws[f"F{r}"] = code, it["description"], it["unit"], it["qty"]
        if s.get("percent"):
            pct_rows.append(r)
            ws[f"I{r}"] = DEFAULT_PCT
            ws[f"I{r}"].number_format = "0%"; ws[f"I{r}"].fill = INPUT_FILL
            ws[f"H{r}"] = f"=ROUNDUP(I{r}*$G$4,0)"
            ws[f"G{r}"] = f"=IF(F{r}>0,H{r}/F{r},H{r})"
            ws[f"J{r}"] = f'=TEXT(I{r},"0%")&" OF TOTAL ITEMS COST"'
        else:
            if s["price"] is not None:
                ws[f"L{r}"], ws[f"M{r}"] = s["year"], s["price"]
            else:
                ws[f"L{r}"].fill = ws[f"M{r}"].fill = INPUT_FILL
                flagged.append((code, it["description"], s["flag"]))
            ws[f"N{r}"] = f'=IF(AND(L{r}<>"",L{r}<$N$6),M{r}*(1+$R$6)^($N$6-L{r}),"")'
            ws[f"O{r}"] = (f'=IF(OR(L{r}="",M{r}=""),"",ROUND(IF(L{r}<$N$6,N{r}*(1+$R$5)^($O$6-$N$6),'
                           f'M{r}*(1+$R$5)^($O$6-L{r})),2))')
            ws[f"G{r}"] = f'=IF(O{r}="","",O{r})'
            ws[f"H{r}"] = f'=IF(G{r}="","",F{r}*G{r})'
            ws[f"K{r}"] = f"=N(H{r})"  # hidden helper: items cost without the % rows (avoids a circular ref)
            notes = []
            if s["flag"]:
                notes.append(s["flag"])
            if s.get("note"):
                notes.append(s["note"])
                checks.append((code, it["description"], s["note"]))
            if rec and rec.get("found") and rec.get("description") and \
                    rec["description"].strip().lower() != it["description"].strip().lower():
                notes.append(f"MDOT: {rec['description']}")
            ws[f"J{r}"] = "; ".join(notes) or None
        # MDOT data block (exact values)
        if rec and rec.get("found"):
            vals = [rec["range"], rec["price_year"], rec["min_qty"], rec["max_qty"], rec["avg_qty"],
                    rec["low_price"], rec["avg_price"], rec["high_price"], s["price"], s["method"]]
        else:
            vals = [rec["range"] if rec else None, None, None, None, None, None, None, None, None,
                    s["method"] if s.get("percent") else "No MDOT bid price found"]
        if s.get("percent") and rec and rec.get("found"):
            vals[-1] = "% of total items cost (MDOT data for reference only)"
        for i, v in enumerate(vals):
            c = ws.cell(r, 20 + i, v); c.fill = DATA_FILL; c.border = BOX
            if 22 <= 20 + i <= 24:
                c.number_format = QTYF
            elif 25 <= 20 + i <= 28:
                c.number_format = MONEY
        for col in list(range(3, 11)) + list(range(12, 16)):
            ws.cell(r, col).border = BOX
        r += 1
    last = r - 1

    # totals
    ws["G4"] = f"=SUM(K{first}:K{last})"
    ws.column_dimensions["K"].hidden = True
    ws["G4"].number_format = MONEY0; ws["G4"].font = BOLD
    sub, cont, tot = r + 1, r + 3, r + 5
    ws[f"C{sub}"] = "SUB-TOTAL"; ws[f"H{sub}"] = f"=SUM(H{first}:H{last})"
    ws[f"D{cont}"] = "Contingencies"; ws[f"E{cont}"] = "LS"; ws[f"F{cont}"] = 1
    ws[f"I{cont}"] = DEFAULT_PCT; ws[f"I{cont}"].number_format = "0%"; ws[f"I{cont}"].fill = INPUT_FILL
    ws[f"H{cont}"] = f"=ROUNDUP(I{cont}*H{sub},0)"; ws[f"G{cont}"] = f"=H{cont}"
    ws[f"J{cont}"] = f'=TEXT(I{cont},"0%")&" OF SUB-TOTAL"'
    ws[f"C{tot}"] = "TOTAL"; ws[f"H{tot}"] = f"=H{sub}+H{cont}"
    for rr in (sub, cont, tot):
        ws[f"C{rr}"].font = ws[f"H{rr}"].font = BOLD
        ws[f"H{rr}"].number_format = MONEY0
        ws[f"G{rr}"].number_format = MONEY0

    # % dropdowns
    dv = DataValidation(type="list", formula1=pct_ref, allow_blank=False,
                        prompt="Pick a percentage, or type your own", promptTitle="Percent",
                        showInputMessage=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    for rr in pct_rows + [cont]:
        dv.add(f"I{rr}")

    # formats
    for rr in range(first, last + 1):
        ws[f"F{rr}"].number_format = QTYF
        ws[f"G{rr}"].number_format = MONEY
        ws[f"H{rr}"].number_format = MONEY
        ws[f"M{rr}"].number_format = MONEY
        ws[f"N{rr}"].number_format = MONEY
        ws[f"O{rr}"].number_format = MONEY
        ws[f"E{rr}"].alignment = Alignment(horizontal="center")
    widths = {"A": 2, "B": 2, "C": 15, "D": 48, "E": 7, "F": 10, "G": 13, "H": 15, "I": 9, "J": 40, "K": 2,
              "L": 8, "M": 13, "N": 12, "O": 12, "P": 2, "Q": 15, "R": 7, "S": 3, "T": 23, "U": 8, "V": 10,
              "W": 10, "X": 10, "Y": 12, "Z": 12, "AA": 13, "AB": 14, "AC": 34}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws.freeze_panes = "E7"
    ws.sheet_view.zoomScale = 85
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"C1:J{tot}"
    return flagged, checks


def build(sq: dict, meta: dict, out: str, refresh=False, log=print):
    codes = [it["item"] for s in sq["sections"] for it in s["items"]]
    prices = mbp.lookup(codes, today=meta["today"], refresh=refresh, log=log)
    downloaded = sorted({p["downloaded"] for p in prices.values() if p.get("found")})
    meta["downloaded"] = ", ".join(downloaded) or meta["today"].isoformat()
    wb = Workbook()
    lists = wb.active; lists.title = "Lists"
    for i in range(1, 21):
        lists.cell(i, 1, i / 100).number_format = "0%"
    lists.sheet_state = "hidden"
    report = {}
    for sec in sq["sections"]:
        report[sec["name"]] = write_sheet(wb, sec, prices, meta, "Lists!$A$1:$A$20")
    wb.move_sheet("Lists", offset=len(wb.sheetnames))
    wb.active = 0
    wb.save(out)
    return prices, report


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("sq_json")
    ap.add_argument("--project", required=True)
    ap.add_argument("--fa", default="", help="federal-aid number, e.g. STBG-9999-05(413)")
    ap.add_argument("--by", default="")
    ap.add_argument("--county")
    ap.add_argument("--out", required=True)
    ap.add_argument("--today")
    ap.add_argument("--refresh", action="store_true")
    a = ap.parse_args()
    sq = json.loads(Path(a.sq_json).read_text())
    today = dt.date.fromisoformat(a.today) if a.today else dt.date.today()
    job = sq.get("job") or ""
    meta = {"project": a.project, "fa": a.fa, "by": a.by.upper(), "county": (a.county or sq.get("county") or "").upper(),
            "today": today, "job_slash": job.replace("-", "/")}
    prices, report = build(sq, meta, a.out, refresh=a.refresh, log=lambda m: print(m, file=sys.stderr))
    print(f"Saved {a.out}")
    for sec, (flagged, checks) in report.items():
        for code, desc, why in flagged:
            print(f"BLANK {sec:14} {code:15} {desc[:45]:45} {why}")
        for code, desc, why in checks:
            print(f"CHECK {sec:14} {code:15} {desc[:45]:45} {why}")


if __name__ == "__main__":
    main()
