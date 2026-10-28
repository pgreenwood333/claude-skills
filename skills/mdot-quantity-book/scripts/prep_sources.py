#!/usr/bin/env python3
"""Turn every calculation file into a clean PDF, ready for the binder.

  prep_sources.py scan <files...> --out <work>              inventory: sheets, print areas, corner numbers, errors
  prep_sources.py convert <work> [--clear "file.xlsx!Sheet!N2" ...] [--landscape ...] [--skip file ...]

scan writes <work>/sources.json, one entry per file (and per worksheet for Excel):
  - kind (excel/word/pdf/text), pages or sheets, print area found vs content area (cells AND floating images,
    so a drawing to the right of the table is not cut off), formula cells whose cached value is an error
    (#NAME?, #REF!, ...) or missing, and CORNER NUMBERS: small integers typed in the top rows at the right
    edge (the engineer's own page numbers). They tell you the order the engineer intended, and they must be
    cleared before printing (only the binder's bottom-center numbers should show).
convert writes <work>/pdf/<file>[__<sheet>].pdf:
  - Excel: formulas flattened to their cached values (plain and array formulas, so LibreOffice cannot turn a
    working cell into #NAME?), print area = content area, fit to one page wide, orientation by shape,
    listed corner cells blanked, hidden sheets skipped, then LibreOffice (soffice) prints it.
  - Word (.docx/.doc): LibreOffice prints it. PDF: copied. Text (.txt): typeset in a monospaced font.
Never edits the user's originals: works on copies in <work>/src.
"""
import argparse, glob, json, os, re, shutil, subprocess, sys

ERR = re.compile(r"^#(NAME\?|REF!|VALUE!|DIV/0!|N/A|NUM!|NULL!|SPILL!|CALC!)")


def soffice():
    for c in ("soffice", "libreoffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        p = shutil.which(c) or (c if os.path.exists(c) else None)
        if p:
            return p
    return None


def kind(path):
    e = os.path.splitext(path)[1].lower()
    return {".xlsx": "excel", ".xlsm": "excel", ".xls": "excel", ".docx": "word", ".doc": "word",
            ".pdf": "pdf", ".txt": "text", ".csv": "text"}.get(e, "other")


def content_area(ws):
    """Rows/cols that hold a value, plus the cells covered by images/drawings."""
    from openpyxl.utils import get_column_letter
    max_r = max_c = 0
    min_r = min_c = 10 ** 9
    for row in ws.iter_rows():
        for c in row:
            if c.value not in (None, ""):
                max_r, max_c = max(max_r, c.row), max(max_c, c.column)
                min_r, min_c = min(min_r, c.row), min(min_c, c.column)
    for img in getattr(ws, "_images", []) + getattr(ws, "_charts", []):
        anc = getattr(img, "anchor", None)
        for part in ("_from", "to"):
            m = getattr(anc, part, None)
            if m is not None:
                max_r, max_c = max(max_r, m.row + 1), max(max_c, m.col + 1)
                min_r, min_c = min(min_r, m.row + 1), min(min_c, m.col + 1)
        ext = getattr(anc, "ext", None)
        if ext is not None and getattr(anc, "_from", None) is not None:   # one-cell anchor: estimate width
            cols = int(ext.width / 9525 / 64) + 1 if getattr(ext, "width", None) else 0
            rows = int(ext.height / 9525 / 20) + 1 if getattr(ext, "height", None) else 0
            max_c, max_r = max(max_c, anc._from.col + 1 + cols), max(max_r, anc._from.row + 1 + rows)
    for m in ws.merged_cells.ranges:
        if m.min_row <= max_r and m.min_col <= max_c:
            max_r, max_c = max(max_r, m.max_row), max(max_c, m.max_col)
    if not max_r:
        return None
    return f"{get_column_letter(min(min_c, 1) if min_c > 1 else 1)}1:{get_column_letter(max_c)}{max_r}"


def corner_numbers(ws, max_col):
    out = []
    for row in ws.iter_rows(min_row=1, max_row=5):
        for c in row:
            if isinstance(c.value, (int, float)) and float(c.value).is_integer() and 1 <= c.value <= 300 \
                    and c.column >= max(3, int(max_col * 0.6)):
                out.append(dict(cell=c.coordinate, value=int(c.value)))
    return out


def scan(files, work):
    import openpyxl
    os.makedirs(os.path.join(work, "src"), exist_ok=True)
    items = []
    for f in files:
        k = kind(f)
        dst = os.path.join(work, "src", os.path.basename(f))
        shutil.copy2(f, dst)
        it = dict(file=os.path.basename(f), kind=k, path=dst)
        if k == "excel" and not f.lower().endswith(".xls"):
            wb_f = openpyxl.load_workbook(dst, data_only=False, keep_vba=f.lower().endswith(".xlsm"))
            wb_v = openpyxl.load_workbook(dst, data_only=True)
            sheets = []
            for ws in wb_f.worksheets:
                wv = wb_v[ws.title]
                area = content_area(ws)
                errs, missing = [], 0
                for row in ws.iter_rows():
                    for c in row:
                        isf = (isinstance(c.value, str) and c.value.startswith("=")) or \
                              type(c.value).__name__ == "ArrayFormula"
                        if isf:
                            v = wv[c.coordinate].value
                            if v is None:
                                missing += 1
                            elif isinstance(v, str) and ERR.match(v):
                                errs.append(f"{c.coordinate}={v}")
                mc = ws.max_column
                sheets.append(dict(sheet=ws.title, hidden=ws.sheet_state != "visible", content_area=area,
                                   print_area=str(ws.print_area) if ws.print_area else None,
                                   corner_numbers=corner_numbers(ws, mc), error_values=errs[:20],
                                   formulas_without_cached_value=missing))
            it["sheets"] = sheets
        elif k == "pdf":
            import pymupdf
            it["pages"] = pymupdf.open(dst).page_count
        items.append(it)
    json.dump(dict(items=items), open(os.path.join(work, "sources.json"), "w"), indent=1)
    for it in items:
        print(f"{it['file']}  [{it['kind']}]")
        for s in it.get("sheets", []):
            flag = []
            if s["hidden"]:
                flag.append("hidden")
            if s["corner_numbers"]:
                flag.append("corner #" + ",".join(f"{c['cell']}={c['value']}" for c in s["corner_numbers"]))
            if s["error_values"]:
                flag.append("ERRORS " + " ".join(s["error_values"][:4]))
            if s["formulas_without_cached_value"]:
                flag.append(f"{s['formulas_without_cached_value']} formulas never calculated")
            print(f"   - {s['sheet']}: area {s['content_area']}  {'; '.join(flag)}")


def flatten_and_setup(path, clear, landscape_sheets):
    import openpyxl
    from openpyxl.worksheet.properties import PageSetupProperties
    keep_vba = path.lower().endswith(".xlsm")
    wb = openpyxl.load_workbook(path, keep_vba=keep_vba)
    wv = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        v = wv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                isf = (isinstance(c.value, str) and c.value.startswith("=")) or \
                      type(c.value).__name__ == "ArrayFormula"
                if isf and v[c.coordinate].value is not None:
                    c.value = v[c.coordinate].value
        for cell in clear.get(ws.title, []):
            ws[cell].value = None
        area = content_area(ws)
        if area:
            ws.print_area = area
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        from openpyxl.utils import range_boundaries
        if area:
            c0, r0, c1, r1 = range_boundaries(area)
            wide = sum((ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width or 9) for c in range(c0, c1 + 1))
            tall = sum((ws.row_dimensions[r].height or 15) / 5.5 for r in range(r0, r1 + 1))
            land = ws.title in landscape_sheets or wide > tall * 0.9
            ws.page_setup.orientation = "landscape" if land else "portrait"
        ws.oddHeader.right.text = ws.oddHeader.right.text or ""
    out = path.replace(".xlsm", ".flat.xlsx") if keep_vba else path.replace(".xlsx", ".flat.xlsx")
    wb.save(out)
    return out, [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]


def to_pdf(path, outdir):
    so = soffice()
    if not so:
        sys.exit("LibreOffice (soffice) is not installed here. In Cowork it is available; on a Mac install "
                 "LibreOffice or export the files to PDF from Excel/Word and pass the PDFs instead.")
    subprocess.run([so, "--headless", "--convert-to", "pdf", "--outdir", outdir, path], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    return os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".pdf")


def text_pdf(path, out):
    import pymupdf
    doc = pymupdf.open()
    lines = open(path, errors="replace").read().expandtabs(4).splitlines()
    width = max([len(l) for l in lines] + [40])
    landscape = width > 95
    rect = pymupdf.paper_rect("letter-l" if landscape else "letter")
    fs = min(10, (rect.width - 72) / (width * 0.6))
    per = int((rect.height - 72) / (fs * 1.25))
    for i in range(0, max(len(lines), 1), per):
        p = doc.new_page(width=rect.width, height=rect.height)
        p.insert_text((36, 36 + fs), "\n".join(lines[i:i + per]), fontname="cour", fontsize=fs, lineheight=1.25)
    doc.save(out)


def convert(work, clears, landscape, skip):
    import openpyxl, pymupdf
    data = json.load(open(os.path.join(work, "sources.json")))
    pdfdir = os.path.join(work, "pdf"); os.makedirs(pdfdir, exist_ok=True)
    clear = {}
    for c in clears:
        f, s, cell = c.split("!")
        clear.setdefault(f, {}).setdefault(s, []).append(cell)
    made = []
    for it in data["items"]:
        if it["file"] in skip:
            continue
        stem = os.path.splitext(it["file"])[0]
        if it["kind"] == "excel":
            flat, visible = flatten_and_setup(it["path"], clear.get(it["file"], {}), set(landscape))
            # print each visible sheet on its own so pages can be ordered by sheet
            for sname in visible:
                wb = openpyxl.load_workbook(flat)
                for ws in wb.worksheets:
                    if ws.title != sname:
                        ws.sheet_state = "hidden"
                wb.active = wb.sheetnames.index(sname)
                one = os.path.join(work, "src", f"{stem}__{re.sub(r'[^A-Za-z0-9]+', '_', sname)}.xlsx")
                wb.save(one)
                pdf = to_pdf(one, pdfdir)
                made.append(dict(file=it["file"], sheet=sname, pdf=pdf, pages=pymupdf.open(pdf).page_count))
        elif it["kind"] == "word":
            pdf = to_pdf(it["path"], pdfdir)
            made.append(dict(file=it["file"], pdf=pdf, pages=pymupdf.open(pdf).page_count))
        elif it["kind"] == "pdf":
            pdf = os.path.join(pdfdir, it["file"]); shutil.copy2(it["path"], pdf)
            made.append(dict(file=it["file"], pdf=pdf, pages=pymupdf.open(pdf).page_count))
        elif it["kind"] == "text":
            pdf = os.path.join(pdfdir, stem + ".pdf"); text_pdf(it["path"], pdf)
            made.append(dict(file=it["file"], pdf=pdf, pages=pymupdf.open(pdf).page_count))
    json.dump(dict(pdfs=made), open(os.path.join(work, "pdfs.json"), "w"), indent=1)
    for m in made:
        print(f"{m['pages']:>2} p  {m['file']}{' / ' + m['sheet'] if m.get('sheet') else ''}  ->  {m['pdf']}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("scan"); s.add_argument("files", nargs="+"); s.add_argument("--out", required=True)
    c = sub.add_parser("convert"); c.add_argument("work")
    c.add_argument("--clear", nargs="*", default=[]); c.add_argument("--landscape", nargs="*", default=[])
    c.add_argument("--skip", nargs="*", default=[])
    a = ap.parse_args()
    if a.cmd == "scan":
        files = [f for pat in a.files for f in (glob.glob(pat) or [pat])]
        scan(files, a.out)
    else:
        convert(a.work, a.clear, a.landscape, a.skip)


if __name__ == "__main__":
    main()
